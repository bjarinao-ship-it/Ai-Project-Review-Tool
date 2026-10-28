import os
import pymupdf
import docx

def extract_text_from_file(file_path: str) -> str:
    """
    Extracts plain text from PDF, DOCX, or TXT file.
    Uses PyMuPDF for PDF, python-docx for DOCX, and standard reading for TXT.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return _extract_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return _extract_from_docx(file_path)
    elif ext in (".txt", ".md", ".json", ".csv"):
        return _extract_from_txt(file_path)
    else:
        # Attempt plain text read as fallback
        try:
            return _extract_from_txt(file_path)
        except Exception:
            raise ValueError(f"Unsupported document format: {ext}")

def _extract_from_pdf(file_path: str) -> str:
    doc = pymupdf.open(file_path)
    text_parts = []
    for page_idx in range(len(doc)):
        page = doc[page_idx]
        page_text = page.get_text("text")
        if page_text:
            text_parts.append(page_text.strip())
    doc.close()
    return "\n\n".join(text_parts)

def _extract_from_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    text_parts = []
    
    # Extract paragraphs
    for p in doc.paragraphs:
        if p.text and p.text.strip():
            text_parts.append(p.text.strip())

    # Extract tables content if any
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_data:
                text_parts.append(" | ".join(row_data))

    return "\n\n".join(text_parts)

def _extract_from_txt(file_path: str) -> str:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            return f.read()

